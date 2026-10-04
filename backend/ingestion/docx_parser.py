from pathlib import Path

from docx import Document

from backend.domain.errors import DocumentParseError
from backend.domain.facts import ParsedDocument
from backend.ingestion.document_type import detect_document_type


def parse_docx(path: str | Path, manifest_type: str = "") -> ParsedDocument:
    p = Path(path)
    try:
        doc = Document(p)
        fields: dict[str, str] = {}
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if len(cells) >= 2:
                    fields[cells[0].rstrip(": ").strip()] = cells[1].strip()
        paragraphs = [paragraph.text.strip() for paragraph in doc.paragraphs if paragraph.text.strip()]
        sections: dict[str, str] = {}
        active = "Document"
        for text in paragraphs:
            if text.lower() in {"assessment", "findings", "impression", "plan", "hospital course",
                                "discharge plan", "procedure / investigation", "procedure", "presenting complaint"}:
                active = text
                sections[active] = ""
            else:
                sections[active] = (sections.get(active, "") + " " + text).strip()
        return ParsedDocument(p.stem, detect_document_type(p, manifest_type), p, fields, sections, paragraphs)
    except Exception as exc:
        raise DocumentParseError(p.name) from exc
